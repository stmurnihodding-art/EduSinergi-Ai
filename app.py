import streamlit as st
from google import genai
from google.genai import types
import io
import time
import docx
from pptx import Presentation
from pptx.util import Pt
from fpdf import FPDF

# 1. Konfigurasi Halaman Dashboard Responsif
st.set_page_config(
    page_title="SEKOLAHKITA AI - Studio Kreasi & Visual",
    page_icon="⚡",
    layout="wide",
    initial_sidebar_state="collapsed"
)

# 2. Styling CSS Responsif & Grid Presisi
custom_css = """
<style>
@import url('https://fonts.googleapis.com/css2?family=Plus+Jakarta+Sans:wght@400;500;600;700;800&family=Syne:wght@700;800&family=Orbitron:wght@800;900&display=swap');

html, body, [class*="css"] {
    font-family: 'Plus Jakarta Sans', sans-serif;
}

/* Header & Judul */
.hero-wrapper {
    display: flex;
    align-items: center;
    gap: 8px;
    margin-bottom: 0px;
    flex-wrap: wrap;
}

.hero-icon {
    font-size: 2rem;
    filter: drop-shadow(0 0 10px rgba(0, 242, 254, 0.7));
}

.title-sekolah {
    font-family: 'Syne', 'Orbitron', sans-serif;
    font-size: 2.1rem;
    font-weight: 900;
    color: #00e5ff;
    letter-spacing: 1px;
    text-shadow: 0 0 16px rgba(0, 229, 255, 0.7);
}

.title-kita {
    font-family: 'Syne', 'Orbitron', sans-serif;
    font-size: 2.1rem;
    font-weight: 900;
    color: #ff7b00;
    letter-spacing: 1px;
    text-shadow: 0 0 16px rgba(255, 123, 0, 0.7);
}

.title-ai {
    font-family: 'Orbitron', sans-serif;
    font-size: 2.1rem;
    font-weight: 900;
    color: #ffffff;
    letter-spacing: 2px;
    text-shadow: 0 0 12px rgba(255, 255, 255, 0.6);
}

/* Neon Glow Bar */
.neon-glow-bar {
    height: 3px;
    width: 100%;
    background: linear-gradient(90deg, #00f2fe 0%, #4facfe 35%, #ff7b00 70%, #ffae19 100%);
    border-radius: 4px;
    box-shadow: 0 0 10px rgba(0, 242, 254, 0.6);
    margin-top: 4px;
    margin-bottom: 14px;
}

/* Tipografi Judul */
.studio-heading {
    font-family: 'Syne', sans-serif;
    font-size: 1.15rem;
    font-weight: 800;
    letter-spacing: 0.8px;
    background: linear-gradient(135deg, #00f2fe 0%, #ffffff 70%);
    -webkit-background-clip: text;
    -webkit-text-fill-color: transparent;
    margin-bottom: 6px;
}

.canvas-heading {
    font-family: 'Syne', sans-serif;
    font-size: 1.35rem;
    font-weight: 800;
    letter-spacing: 0.8px;
    background: linear-gradient(135deg, #ffffff 0%, #ffae19 50%, #ff7b00 100%);
    -webkit-background-clip: text;
    -webkit-text-fill-color: transparent;
    margin-top: 10px;
    margin-bottom: 8px;
}

/* Kotak Border Kontainer */
div[data-testid="stVerticalBlockBorderWrapper"] {
    border-radius: 10px !important;
    border: 1px solid rgba(0, 242, 254, 0.35) !important;
    box-shadow: 0 0 14px rgba(0, 242, 254, 0.08) !important;
    background-color: rgba(255, 255, 255, 0.02) !important;
    padding: 8px 10px !important;
}

/* Gambar Kartu Mini Presisi */
div[data-testid="stImage"] img {
    height: 68px !important;
    width: 100% !important;
    object-fit: cover !important;
    border-radius: 6px !important;
}

div[data-testid="stCaptionContainer"] {
    text-align: center !important;
    font-weight: 600 !important;
    color: #cbd5e1 !important;
    font-size: 0.72rem !important;
    margin-top: 2px !important;
    white-space: nowrap !important;
    overflow: hidden !important;
    text-overflow: ellipsis !important;
}

/* Tombol Kartu Neon */
.btn-cyan > button {
    font-size: 0.74rem !important;
    font-weight: 700 !important;
    border-radius: 6px !important;
    background: rgba(0, 242, 254, 0.08) !important;
    border: 1px solid #00f2fe !important;
    color: #00f2fe !important;
    box-shadow: 0 0 8px rgba(0, 242, 254, 0.3) !important;
    width: 100% !important;
    padding: 3px 6px !important;
}

.btn-emerald > button {
    font-size: 0.74rem !important;
    font-weight: 700 !important;
    border-radius: 6px !important;
    background: rgba(16, 185, 129, 0.08) !important;
    border: 1px solid #10b981 !important;
    color: #10b981 !important;
    box-shadow: 0 0 8px rgba(16, 185, 129, 0.3) !important;
    width: 100% !important;
    padding: 3px 6px !important;
}

.btn-purple > button {
    font-size: 0.74rem !important;
    font-weight: 700 !important;
    border-radius: 6px !important;
    background: rgba(168, 85, 247, 0.08) !important;
    border: 1px solid #a855f7 !important;
    color: #c084fc !important;
    box-shadow: 0 0 8px rgba(168, 85, 247, 0.3) !important;
    width: 100% !important;
    padding: 3px 6px !important;
}

.btn-pink > button {
    font-size: 0.74rem !important;
    font-weight: 700 !important;
    border-radius: 6px !important;
    background: rgba(236, 72, 153, 0.08) !important;
    border: 1px solid #ec4899 !important;
    color: #f472b6 !important;
    box-shadow: 0 0 8px rgba(236, 72, 153, 0.3) !important;
    width: 100% !important;
    padding: 3px 6px !important;
}

/* Tombol Utama Generate */
.main-btn > button {
    font-family: 'Plus Jakarta Sans', sans-serif !important;
    font-weight: 700 !important;
    font-size: 0.95rem !important;
    border-radius: 8px !important;
    background: linear-gradient(90deg, #ff7b00, #e65c00) !important;
    color: white !important;
    border: none !important;
    box-shadow: 0 0 14px rgba(255, 123, 0, 0.45) !important;
    margin-top: 4px;
}
.main-btn > button:hover {
    box-shadow: 0 0 22px rgba(255, 123, 0, 0.8) !important;
    transform: translateY(-1px);
}

/* Tombol Revisi */
.refine-btn > button {
    font-family: 'Plus Jakarta Sans', sans-serif !important;
    font-weight: 700 !important;
    border-radius: 8px !important;
    background: linear-gradient(90deg, #00e5ff, #0099cc) !important;
    color: black !important;
    border: none !important;
    box-shadow: 0 0 12px rgba(0, 229, 255, 0.4) !important;
}
.refine-btn > button:hover {
    box-shadow: 0 0 20px rgba(0, 229, 255, 0.8) !important;
    color: white !important;
}

/* Styling Tombol Unduh */
div[data-testid="stDownloadButton"] > button {
    font-weight: 600 !important;
    font-size: 0.82rem !important;
    border-radius: 8px !important;
    border: 1px solid rgba(0, 242, 254, 0.5) !important;
    background: rgba(0, 242, 254, 0.08) !important;
    color: #00f2fe !important;
    transition: all 0.2s ease !important;
}
div[data-testid="stDownloadButton"] > button:hover {
    background: #00f2fe !important;
    color: #000000 !important;
    box-shadow: 0 0 14px rgba(0, 242, 254, 0.7) !important;
}

@media (max-width: 768px) {
    .title-sekolah, .title-kita, .title-ai {
        font-size: 1.5rem !important;
    }
    .hero-icon {
        font-size: 1.5rem !important;
    }
    div[data-testid="stImage"] img {
        height: 65px !important;
    }
    .main-btn > button, .btn-cyan > button, .btn-emerald > button, .btn-purple > button, .btn-pink > button {
        min-height: 40px !important;
        font-size: 0.8rem !important;
    }
}
</style>
"""

st.markdown(custom_css, unsafe_allow_html=True)


st.caption("Platform Tata Kelola Sekolah, Modul Ajar, dan Administrasi Terintegrasi (SD - SMP - SMA)")
st.markdown('<div class="neon-glow-bar"></div>', unsafe_allow_html=True)
# 3. Header Judul dengan Logo Resmi JacS
col_logo, col_title = st.columns([0.08, 0.92], gap="small")
with col_logo:
    try:
        st.image("logo.png", width=46)
    except Exception:
        st.markdown("<span class='hero-icon'>⚡</span>", unsafe_allow_html=True)
with col_title:
    st.markdown(
        """
        <div class="hero-wrapper" style="margin-top: 2px;">
            <span class="title-sekolah">JacS</span>
            <span class="title-kita">AI</span>
            <span class="title-ai">GENERATOR</span>
        </div>
        """,
        unsafe_allow_html=True
    )
# Inisialisasi State Formulir
if "selected_role" not in st.session_state:
    st.session_state.selected_role = "Guru Mata Pelajaran (SMP/SMA/SMK)"
if "selected_doc" not in st.session_state:
    st.session_state.selected_doc = ""
if "selected_details" not in st.session_state:
    st.session_state.selected_details = ""
if "generated_doc" not in st.session_state:
    st.session_state.generated_doc = ""
if "generated_image" not in st.session_state:
    st.session_state.generated_image = None
if "uploaded_preview_img" not in st.session_state:
    st.session_state.uploaded_preview_img = None
if "uploaded_preview_doc" not in st.session_state:
    st.session_state.uploaded_preview_doc = ""

# 4. Panel Samping (Sidebar)
with st.sidebar:
    st.header("⚙️ Pengaturan")
    secret_key = st.secrets.get("GEMINI_API_KEY", "")
    if secret_key:
        api_key = secret_key
        st.success("✅ API Key terhubung otomatis")
    else:
        api_key = st.text_input("Masukkan Gemini API Key:", type="password")
        st.markdown("[Dapatkan API Key di Google AI Studio](https://aistudio.google.com/)")

SYSTEM_INSTRUCTION = """
Anda adalah "SEKOLAHKITA AI", asisten komprehensif tata kelola sekolah, perancangan beban mengajar guru (JTM), kurikulum instruksional (SD, SMP, SMA/SMK), tata naskah dinas, slide presentasi materi, dan perancang tata letak visual untuk Canva.
Kemampuan Multi-Format Dokumen & Gambar:
- Jika pengguna mengunggah berkas teks dokumen (PDF, DOCX, TXT), baca dan tindaklanjuti isinya sesuai permintaan.
- Jika pengguna mengunggah gambar (JPG, JPEG, PNG, WEBP), lakukan Optical Character Recognition (OCR) dan analisis secara kontekstual.
- Khusus untuk peran Guru Kelas (SD): kuasai pendekatan tematik, fase fondasi A-C Kurikulum Merdeka, literasi-numerasi dini, serta lembar aktivitas peserta didik (LKPD) ramah anak.
- Jika dokumen berkaitan dengan infografik, poster, LKPD, atau slide presentasi: sertakan poin-poin terstruktur per slide/bagian, kode warna hex, dan panduan transfer ke Canva.
Gunakan bahasa Indonesia baku, formal, dan rapi sesuai tata naskah dinas pendidikan.
Sajikan langsung format dokumen atau matriks tabel siap pakai tanpa basa-basi pembuka.
"""

# Generator Berkas Unduhan
def create_docx_bytes(markdown_text):
    doc = docx.Document()
    doc.add_heading("Draf Hasil Rancangan - SEKOLAHKITA AI", level=1)
    for line in markdown_text.split("\n"):
        clean = line.strip()
        if clean.startswith("### "):
            doc.add_heading(clean.replace("### ", ""), level=3)
        elif clean.startswith("## "):
            doc.add_heading(clean.replace("## ", ""), level=2)
        elif clean.startswith("# "):
            doc.add_heading(clean.replace("# ", ""), level=1)
        elif clean.startswith("- ") or clean.startswith("* "):
            doc.add_paragraph(clean[2:], style='List Bullet')
        elif clean:
            doc.add_paragraph(clean)
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()

def create_pdf_bytes(markdown_text):
    pdf = FPDF()
    pdf.add_page()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.set_font("Helvetica", size=10)
    pdf.cell(0, 10, text="Draf Dokumen - SEKOLAHKITA AI", new_x="LMARGIN", new_y="NEXT", align='C')
    pdf.ln(5)
    for line in markdown_text.split("\n"):
        clean = line.strip()
        safe_line = clean.encode('latin-1', 'replace').decode('latin-1')
        if safe_line:
            pdf.multi_cell(0, 6, text=safe_line)
            pdf.ln(1)
    return bytes(pdf.output())

def create_pptx_bytes(markdown_text):
    prs = Presentation()
    slide_layout = prs.slide_layouts[1]
    lines = [l.strip() for l in markdown_text.split("\n") if l.strip()]
    chunk_size = 5
    for i in range(0, len(lines), chunk_size):
        chunk = lines[i:i+chunk_size]
        slide = prs.slides.add_slide(slide_layout)
        title = slide.shapes.title
        title.text = "SEKOLAHKITA AI - Slide Bahan Ajar"
        body = slide.shapes.placeholders[1]
        tf = body.text_frame
        for idx, item in enumerate(chunk):
            p = tf.add_paragraph() if idx > 0 else tf.paragraphs[0]
            p.text = item
            p.font.size = Pt(14)
    out = io.BytesIO()
    prs.save(out)
    return out.getvalue()

# Fungsi Renderer Hasil, Editor Naskah, Tombol Unduh & Refine Prompt
def render_canvas_content(container):
    with container:
        if st.session_state.uploaded_preview_img:
            with st.expander("📷 **Lihat Foto/Dokumen Asli yang Dibaca AI**", expanded=False):
                st.image(st.session_state.uploaded_preview_img, caption="Dokumen Visual yang Diunggah", use_container_width=True)

        if st.session_state.uploaded_preview_doc:
            with st.expander("📄 **Lihat Kutipan Dokumen Teks/PDF/Word yang Dibaca AI**", expanded=False):
                st.text_area("Isi Teks Dokumen:", value=st.session_state.uploaded_preview_doc, height=120, disabled=True)

        if st.session_state.generated_image:
            st.markdown("##### 🖼️ Hasil Gambar Ilustrasi Sesuai Permintaan:")
            st.image(st.session_state.generated_image, use_container_width=True)
            st.download_button(
                label="💾 Unduh Gambar Ilustrasi (PNG)",
                data=st.session_state.generated_image,
                file_name="sekolahkita_visual.png",
                mime="image/png",
                use_container_width=True
            )
            st.divider()

        if st.session_state.generated_doc:
            st.markdown("##### 📥 Pusat Unduh Berkas Hasil Kreasi:")
            col_d1, col_d2, col_d3, col_d4 = st.columns(4)
            with col_d1:
                docx_bytes = create_docx_bytes(st.session_state.generated_doc)
                st.download_button(
                    label="📄 Unduh Word (.docx)",
                    data=docx_bytes,
                    file_name="Rancangan_SekolahKita_AI.docx",
                    mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                    use_container_width=True
                )
            with col_d2:
                pdf_bytes = create_pdf_bytes(st.session_state.generated_doc)
                st.download_button(
                    label="📑 Unduh PDF (.pdf)",
                    data=pdf_bytes,
                    file_name="Rancangan_SekolahKita_AI.pdf",
                    mime="application/pdf",
                    use_container_width=True
                )
            with col_d3:
                pptx_bytes = create_pptx_bytes(st.session_state.generated_doc)
                st.download_button(
                    label="📊 Unduh Slide (.pptx)",
                    data=pptx_bytes,
                    file_name="Materi_SekolahKita_AI.pptx",
                    mime="application/vnd.openxmlformats-officedocument.presentationml.presentation",
                    use_container_width=True
                )
            with col_d4:
                st.download_button(
                    label="📝 Unduh Teks (.txt)",
                    data=st.session_state.generated_doc,
                    file_name="Rancangan_SekolahKita_AI.txt",
                    mime="text/plain",
                    use_container_width=True
                )

            st.write("")
            tab_view, tab_edit, tab_copy, tab_canva = st.tabs([
                "👁️ Tampilan Naskah & Rancangan",
                "✏️ Edit Naskah Langsung",
                "📋 Salin Format Naskah",
                "🎨 Panduan Buka di Canva"
            ])
            
            with tab_view:
                st.markdown(st.session_state.generated_doc)

            with tab_edit:
                st.caption("Anda dapat mengubah, menambahkan, atau memotong isi naskah secara langsung di bawah ini:")
                edited_text = st.text_area("Editor Teks:", value=st.session_state.generated_doc, height=350, key="editor_naskah_area")
                if st.button("💾 Simpan Perubahan Naskah", use_container_width=True):
                    st.session_state.generated_doc = edited_text
                    st.success("✅ Perubahan naskah berhasil disimpan! File unduhan otomatis diperbarui.")
                    st.rerun()

            with tab_copy:
                st.caption("Klik tombol salin di pojok kanan atas untuk menempelkannya ke lembar kerja Anda:")
                st.code(st.session_state.generated_doc, language="markdown")

            with tab_canva:
                st.markdown("##### 🚀 Lanjutkan Desain ke Canva")
                st.write("Jika Anda memerlukan templat desain visual, tata letak poster, atau LKPD interaktif, klik tautan Canva di bawah ini:")
                st.link_button("🌐 Buka Editor Canva", "https://www.canva.com/")
                st.info("💡 **Tips Guru:** Salin teks atau tabel dari tab *Salin Format Naskah*, lalu tempelkan langsung ke kotak teks atau elemen tabel di dalam Canva.")

            # --- FITUR AJUKAN REVISI / RE-GENERATE DENGAN PERINTAH BARU ---
            st.divider()
            st.markdown("##### 🔄 Ajukan Revisi / Lanjutkan Permintaan:")
            refine_input = st.text_input(
                "Ingin menambah, mengubah, atau memperdalam materi di atas?",
                placeholder="Contoh: 'Tolong tambahkan 5 soal kuis esai', 'Ubah pendekatan materi menjadi lebih santai untuk siswa', dsb...",
                key="input_refine_prompt"
            )
            st.markdown('<div class="refine-btn">', unsafe_allow_html=True)
            btn_refine = st.button("🚀 Kirim Revisi / Generate Ulang", use_container_width=True)
            st.markdown('</div>', unsafe_allow_html=True)

            if btn_refine:
                if not refine_input:
                    st.warning("Silakan ketik instruksi revisi yang diinginkan terlebih dahulu.")
                elif not api_key:
                    st.error("API Key belum terpasang.")
                else:
                    try:
                        client = genai.Client(api_key=api_key)
                        refine_payload = [
                            f"Berikut naskah awal yang sudah dibuat:\n{st.session_state.generated_doc}\n\n"
                            f"Instruksi Revisi/Tambahan dari Pengguna:\n{refine_input}\n\n"
                            f"Tolong sesuaikan, revisi, atau tambahkan naskah tersebut secara lengkap sesuai instruksi revisi:"
                        ]
                        with st.spinner("Sedang memperbarui rancangan sesuai instruksi revisi Anda..."):
                            for model_name in ["gemini-3.6-flash", "gemini-2.5-flash"]:
                                try:
                                    resp = client.models.generate_content(
                                        model=model_name,
                                        contents=refine_payload,
                                        config={"system_instruction": SYSTEM_INSTRUCTION}
                                    )
                                    if resp and resp.text:
                                        st.session_state.generated_doc = resp.text
                                        st.rerun()
                                except Exception as err:
                                    if "503" in str(err):
                                        time.sleep(1)
                                        continue
                                    else:
                                        raise err
                    except Exception as e:
                        st.error(f"Gagal melakukan revisi: {e}")

# 5. Tata Letak 2 Kolom: Studio di Kiri (2.3) - 4 Kartu di Kanan (1.0)
col_left, col_right = st.columns([2.3, 1.0], gap="medium")

# --- SISI KIRI: STUDIO RANCANGAN ---
roles_list = [
    "Guru Kelas (SD)",
    "Guru Mata Pelajaran (SMP/SMA/SMK)",
    "Wali Kelas",
    "Kepala Sekolah",
    "Tim Kurikulum",
    "Tim Kesiswaan",
    "Tata Usaha (TU)"
]
current_role_index = roles_list.index(st.session_state.selected_role) if st.session_state.selected_role in roles_list else 1

with col_left:
    with st.container(border=True):
        st.markdown('<div class="studio-heading">🎛️ STUDIO RANCANGAN</div>', unsafe_allow_html=True)
        role = st.selectbox(
            "Pilih Peran Anda:",
            roles_list,
            index=current_role_index
        )
        doc_type = st.text_input(
            "Jenis Dokumen / Administrasi:",
            value=st.session_state.selected_doc,
            placeholder="Contoh: Modul Ajar Tematik, Catatan Rapor, Slide Pembelajaran"
        )
        details = st.text_area(
            "Detail Tambahan / Instruksi Tindak Lanjut:",
            value=st.session_state.selected_details,
            placeholder="Ketik instruksi tindak lanjut (misal: 'buatkan materi slide PPT', 'analisis foto SK ini', atau 'buatkan gambar visual')...",
            height=100
        )
        
        uploaded_file = st.file_uploader(
            "📎 Unggah Dokumen / Foto Berkas untuk Dibaca AI:",
            type=["docx", "pdf", "txt", "jpg", "jpeg", "png", "webp"],
            help="Unggah berkas Word (.docx), PDF, teks (.txt), atau foto dokumen (.jpg/.png) untuk dianalisis dan ditindaklanjuti secara otomatis."
        )

        st.markdown('<div class="main-btn">', unsafe_allow_html=True)
        btn_generate = st.button("⚡ Hasilkan Rancangan (Generate)", use_container_width=True)
        st.markdown('</div>', unsafe_allow_html=True)

# --- SISI KANAN: 4 KARTU GRID 2x2 KOMPAK ---
with col_right:
    r1_col1, r1_col2 = st.columns(2, gap="small")
    with r1_col1:
        with st.container(border=True):
            st.image("https://images.unsplash.com/photo-1580582932707-520aed937b7b?w=400&q=80", caption="Modul Ajar", use_container_width=True)
            st.markdown('<div class="btn-cyan">', unsafe_allow_html=True)
            if st.button("📘 Modul", key="btn_tpl_1", use_container_width=True):
                st.session_state.selected_role = "Guru Mata Pelajaran (SMP/SMA/SMK)"
                st.session_state.selected_doc = "Modul Ajar Kurikulum Merdeka"
                st.session_state.selected_details = "Rancang modul ajar tematik/mapel lengkap: Fase/Kelas, Capaian Pembelajaran (CP), Tujuan Pembelajaran (TP), aktivitas berdiferensiasi, LKPD anak, dan konsep tata letak Canva."
                st.rerun()
            st.markdown('</div>', unsafe_allow_html=True)

    with r1_col2:
        with st.container(border=True):
            st.image("https://images.unsplash.com/photo-1454165804606-c3d57bc86b40?w=400&q=80", caption="Asesmen Rapor", use_container_width=True)
            st.markdown('<div class="btn-emerald">', unsafe_allow_html=True)
            if st.button("📝 Rapor", key="btn_tpl_2", use_container_width=True):
                st.session_state.selected_role = "Wali Kelas"
                st.session_state.selected_doc = "Catatan Rapor Siswa"
                st.session_state.selected_details = "Kompilasi narasi catatan rapor yang konstruktif dan memotivasi untuk siswa (akademik tinggi, butuh bimbingan membaca/hitung, penguatan karakter Profil Pelajar Pancasila)."
                st.rerun()
            st.markdown('</div>', unsafe_allow_html=True)

    r2_col1, r2_col2 = st.columns(2, gap="small")
    with r2_col1:
        with st.container(border=True):
            st.image("https://images.unsplash.com/photo-1450133064473-71024230f91b?w=400&q=80", caption="Regulasi SK", use_container_width=True)
            st.markdown('<div class="btn-purple">', unsafe_allow_html=True)
            if st.button("🏛️ SK Dinas", key="btn_tpl_3", use_container_width=True):
                st.session_state.selected_role = "Kepala Sekolah"
                st.session_state.selected_doc = "Surat Keputusan (SK) Beban Kerja Guru"
                st.session_state.selected_details = "Draf naskah dinas resmi SK Kepala Sekolah tentang Pembagian Tugas Guru Kelas/Mapel Tahun Ajaran Baru, konsideran menimbang/mengingat, serta lampiran rincian beban tugas."
                st.rerun()
            st.markdown('</div>', unsafe_allow_html=True)

    with r2_col2:
        with st.container(border=True):
            st.image("https://images.unsplash.com/photo-1577896851231-70ef18881754?w=400&q=80", caption="Kesiswaan", use_container_width=True)
            st.markdown('<div class="btn-pink">', unsafe_allow_html=True)
            if st.button("👥 Disiplin", key="btn_tpl_4", use_container_width=True):
                st.session_state.selected_role = "Tim Kesiswaan"
                st.session_state.selected_doc = "Program Pembiasaan Karakter & Tata Tertib"
                st.session_state.selected_details = "Pedoman pembiasaan budaya positif dan tata tertib siswa, panduan kegiatan sekolah ramah anak, serta jadwal MPLS."
                st.rerun()
            st.markdown('</div>', unsafe_allow_html=True)

# --- STUDIO KREASI & VISUAL ---
st.write("")
st.markdown('<div class="canvas-heading">🎨 STUDIO KREASI & VISUAL</div>', unsafe_allow_html=True)
canvas_box = st.container(border=True)

# Logika Pemrosesan Berkas Dokumen dengan Fallback & Auto-Retry
if btn_generate:
    if not api_key:
        st.error("Silakan masukkan Gemini API Key di menu samping terlebih dahulu.")
    elif not doc_type and not uploaded_file:
        st.warning("Mohon sebutkan jenis dokumen atau unggah berkas yang ingin diproses.")
    else:
        try:
            client = genai.Client(api_key=api_key)
            combined_input = f"{doc_type} {details}".lower()
            keywords_image = ["buatkan gambar", "hasilkan gambar", "generate image", "ilustrasi gambar", "buat gambar"]
            need_real_image = any(kw in combined_input for kw in keywords_image)

            st.session_state.generated_image = None
            st.session_state.uploaded_preview_img = None
            st.session_state.uploaded_preview_doc = ""

            target_doc_title = doc_type if doc_type else "Analisis dan Tindak Lanjut Berkas Sumber"
            text_prompt = f"Peran Pengguna: {role}\nJenis Dokumen/Target: {target_doc_title}\nInstruksi Tambahan: {details}\n\nLakukan analisis menyeluruh dan tindak lanjuti dokumen/materi ini secara terstruktur:"
            
            contents_payload = [text_prompt]

            # Ekstraksi Berkas
            if uploaded_file is not None:
                file_bytes = uploaded_file.getvalue()
                file_name = uploaded_file.name.lower()

                if file_name.endswith(('.jpg', '.jpeg', '.png', '.webp')):
                    st.session_state.uploaded_preview_img = file_bytes
                    mime = uploaded_file.type if uploaded_file.type else "image/jpeg"
                    contents_payload.append(types.Part.from_bytes(data=file_bytes, mime_type=mime))

                elif file_name.endswith('.txt'):
                    txt_content = file_bytes.decode('utf-8', errors='ignore')
                    st.session_state.uploaded_preview_doc = txt_content[:2000]
                    contents_payload.append(f"\n--- ISI DOKUMEN TEKS ({uploaded_file.name}) ---\n{txt_content}")

                elif file_name.endswith('.docx'):
                    try:
                        doc = docx.Document(io.BytesIO(file_bytes))
                        docx_text = "\n".join([p.text for p in doc.paragraphs if p.text])
                    except Exception:
                        docx_text = file_bytes.decode('latin-1', errors='ignore')
                    st.session_state.uploaded_preview_doc = docx_text[:2000]
                    contents_payload.append(f"\n--- ISI DOKUMEN WORD ({uploaded_file.name}) ---\n{docx_text}")

                elif file_name.endswith('.pdf'):
                    try:
                        import pypdf
                        reader = pypdf.PdfReader(io.BytesIO(file_bytes))
                        pdf_text = "\n".join([page.extract_text() or "" for page in reader.pages])
                    except Exception:
                        pdf_text = file_bytes.decode('latin-1', errors='ignore')
                    st.session_state.uploaded_preview_doc = pdf_text[:2000]
                    contents_payload.append(f"\n--- ISI DOKUMEN PDF ({uploaded_file.name}) ---\n{pdf_text}")

            candidate_models = ["gemini-3.6-flash", "gemini-2.5-flash"]
            response_text = None

            with st.spinner("Sedang menyusun naskah dan rancangan materi..."):
                for model_name in candidate_models:
                    try:
                        resp = client.models.generate_content(
                            model=model_name,
                            contents=contents_payload,
                            config={"system_instruction": SYSTEM_INSTRUCTION}
                        )
                        if resp and resp.text:
                            response_text = resp.text
                            break
                    except Exception as err:
                        if "503" in str(err):
                            time.sleep(1)
                            continue
                        else:
                            raise err

            if response_text:
                st.session_state.generated_doc = response_text
            else:
                st.error("Layanan AI sedang sibuk sementara. Silakan klik kembali tombol 'Hasilkan Rancangan (Generate)'.")

            # Pembuatan Gambar Ilustrasi jika Diminta
            if need_real_image and response_text:
                with st.spinner("Sedang merender ilustrasi visual..."):
                    try:
                        img_response = client.models.generate_images(
                            model='imagen-3.0-generate-002',
                            prompt=f"Educational illustration for school, classroom context: {target_doc_title}, {details}",
                            config=types.GenerateImagesConfig(
                                number_of_images=1,
                                aspect_ratio="16:9"
                            )
                        )
                        for generated_image in img_response.generated_images:
                            st.session_state.generated_image = generated_image.image.image_bytes
                    except Exception:
                        pass

            render_canvas_content(canvas_box)

        except Exception as e:
            st.error(f"Terjadi kesalahan saat memproses: {e}")
else:
    if st.session_state.generated_doc or st.session_state.generated_image:
        render_canvas_content(canvas_box)
    else:
        with canvas_box:
            st.info("Draf dokumen dinas, matriks tabel, slide materi, atau konsep visual Canva akan tampil di lembar ini.")
